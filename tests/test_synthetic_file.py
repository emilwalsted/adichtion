import io
import os
import tempfile

import numpy as np
import pytest

from adichtion import AdichtFile, AdichtContainer
from adichtion.export import read_labchart_text, write_text

LSB_FLOW = 0.2 * 1.024 / 32768
LSB_10V = 10.0 * 1.024 / 32768


def test_container_tree(synthetic):
    with AdichtContainer(synthetic["path"]) as c:
        names = sorted(e.path for e in c.walk())
        assert names == ["Blocks_0", "ChildDocuments", "ChildDocuments/ADIStdPart4", "Map", "Version"]
        assert c.read_stream(c.find_stream("Version")).startswith(b"\x01\x00\x00\x80")


def test_parse_mode(any_document):
    with AdichtFile(any_document["path"]) as f:
        assert f.parse_mode == ("scan" if any_document["layout"] == "legacy" else "sequential")
        assert f.n_channels == 4


def test_metadata(any_document):
    with AdichtFile(any_document["path"]) as f:
        assert f.channel_titles == ["Flow", "Pmo", "Calc", "Pdi"]
        assert [c.dtype for c in f.channels] == ["int16", "int16", "float32", "float32"]
        assert [c.range_volts for c in f.channels][:2] == [0.2, 10.0]
        assert f.channels[1].units.units == "cmH2O"
        assert f.channels[1].units.enabled
        assert f.channels[0].units is None or not f.channels[0].units.enabled
        assert [r.ticks for r in f.records] == [40000, 5000]
        assert [r.block_ticks for r in f.records] == [[16000, 16000, 8000], [5000]]
        r = f.records[0]
        assert [(c.tick, c.channel, c.text) for c in r.comments] == [(10, -1, "start"), (16005, 1, "15,5 inH2O"), (39995, 0, "stop")]
        assert r.selection_between("start", "stop").ticks == 39985
        assert r.sample_rate_hz(2) == 1000.0
        assert r.start_local.utcoffset().total_seconds() == 3600
        calc_label = "" if any_document["layout"] == "legacy" else "V"   # the legacy layout carries no LSB
        assert [r.value_label(c) for c in range(4)] == ["V", "cmH2O", calc_label, "cmH2O"]


def test_channel_entries(documents):
    for name in ("v29", "v26"):
        with AdichtFile(documents[name]["path"]) as f:
            for rec in f.records:
                for ch in range(4):
                    cr = f.channel_record(rec.index, ch)
                    assert cr is not None and cr.channel == ch and cr.ticks == rec.ticks
                    assert [b.ticks for b in cr.blocks] == rec.block_ticks
                    # the block index points at the sub-blocks the walk finds
                    for b, blk in enumerate(cr.blocks):
                        assert blk.offset == f.index.get(rec.index, b, ch).offset
                    # later records reference the units object of the first record
                    assert cr.units is f.channel_record(0, ch).units
            assert f.channel_record(0, 1).units.lsb == pytest.approx(LSB_10V)
            assert f.channel_record(0, 2).units.lsb == pytest.approx(LSB_FLOW)
            if name == "v29":
                assert f.chart.setups[0].range_volts == 0.2
                assert f.channel_record(1, 0).setup is f.channel_record(0, 0).setup
            else:
                assert f.chart.setups == {}
            walk = f.walk_subblocks()
            assert [s.offset for s in walk] == [s.offset for s in f.index.subblocks]


def test_samples_exact(any_document):
    with AdichtFile(any_document["path"]) as f:
        for ri, rec in enumerate(f.records):
            ref = any_document["records"][ri].data
            assert np.array_equal(rec.read_raw(0), ref[0])
            assert np.array_equal(rec.read_raw(1), ref[1])
            assert np.allclose(rec.read_raw(2), ref[2], rtol=1e-6, atol=1e-3)
            assert np.allclose(rec.read_raw(3), ref[3], rtol=1e-6, atol=1e-1)
            # partial reads across block boundaries
            assert np.array_equal(rec.read_raw(0, 15990, 16010), ref[0][15990:16010])
            # units conversion: counts -> volts -> cmH2O
            volts = ref[1][:5] * LSB_10V
            assert np.allclose(rec.read(1, 0, 5), 13.34 * volts - 0.635)
            assert np.allclose(rec.read(0, 0, 5), ref[0][:5] * LSB_FLOW)


def test_float32_scaling(documents):
    """Calculated channels are scaled by their LSB and converted when the conversion is on."""
    for name in ("v29", "v26"):
        with AdichtFile(documents[name]["path"]) as f:
            rec = f.records[0]
            ref = documents[name]["records"][0].data
            assert np.allclose(rec.read(2, 0, 100), ref[2][:100] * LSB_FLOW, rtol=1e-6)
            assert np.allclose(rec.read(3, 0, 100), ref[3][:100] * 2.0 ** -15 * 1.024, rtol=1e-6)
            assert rec.value_label(2) == "V" and rec.value_label(3) == "cmH2O"


def test_block_envelope(synthetic):
    with AdichtFile(synthetic["path"]) as f:
        rec = f.records[0]
        ref = synthetic["records"][0].data
        t0, t1, lo, hi = rec.block_envelope(0, units=False)
        assert np.allclose(t0, [0, 16, 32]) and np.allclose(t1, [16, 32, 40])
        assert lo[0] == ref[0][:16000].min() and hi[2] == ref[0][32000:].max()
        t0, t1, lo, hi = rec.block_envelope(1)
        assert np.allclose(lo, 13.34 * np.array([ref[1][:16000].min(), ref[1][16000:32000].min(), ref[1][32000:].min()]) * LSB_10V - 0.635)
        # a coarse overview comes from the index, a fine one from the samples; both bracket the data
        t, mn, mx = rec.overview(0, 2, units=False)
        assert len(t) == 2 and mn[0] <= ref[0][:20000].min() and mx[0] >= ref[0][:20000].max()
        t, mn, mx = rec.overview(0, 100)
        assert len(t) == len(mn) == len(mx) == 100
        assert np.all(mn <= mx)
        assert np.allclose(rec.time(0, 3), [0, 0.001, 0.002])


def test_videos(synthetic):
    with AdichtFile(synthetic["path"]) as f:
        assert [v.file_name for v in f.videos] == ["Record1.wmv", ""]
        v = f.records[0].video
        assert v.has_video and v.offset_s == 0.46 and v.wall_duration_s == pytest.approx(38.98)
        assert v.delay_s == pytest.approx(0.17) and v.file_time(2.5) == pytest.approx(2.96)
        assert not f.records[1].video.has_video
        assert f.movies_folder.endswith("synthetic_v29 movies")
        assert f.records[0].video_file is None            # the folder does not exist
        os.makedirs(f.movies_folder, exist_ok=True)
        open(os.path.join(f.movies_folder, "Record1.wmv"), "wb").close()
        assert f.records[0].video_file.endswith("Record1.wmv")


def test_export_roundtrip(synthetic):
    with AdichtFile(synthetic["path"]) as f:
        rec = f.records[0]
        sel = rec.selection(15.99, 16.02)
        buf = io.StringIO()
        n = write_text(rec, sel, [0, 1], buf, fmt="txt", decimal=",")
        assert n == 30
        text = buf.getvalue()
        assert text.startswith("Interval=\t0,001 s\n")
        assert "#* 15,5 inH2O" in text
        p = os.path.join(tempfile.mkdtemp(), "x.txt")
        open(p, "w").write(text)
        header, t, vals, comments = read_labchart_text(p)
        assert header["ChannelTitle"] == "Flow\tPmo"
        assert np.allclose(t[:2], [15.99, 15.991])
        assert np.allclose(vals[:, 1], rec.read(1, sel.start, sel.stop), rtol=1e-6)
        assert comments == [(15, "15,5 inH2O")]


def test_export_whole_file_txt_csv_xlsx(synthetic, tmp_path):
    from adichtion.export import XLSX_MAX_ROWS, export_selections, whole_file
    with AdichtFile(synthetic["path"]) as f:
        sels = whole_file(f)
        assert [(s.record, s.start, s.stop) for s in sels] == [(0, 0, 40000), (1, 0, 5000)]
        # txt: the header lines are repeated at the start of every record
        p = tmp_path / "all.txt"
        n = export_selections(f, sels, [0, 1], p, fmt="txt")
        assert n == 45000
        text = p.read_text()
        assert text.count("Interval=\t") == 2 and text.count("ChannelTitle=\t") == 2
        header, t, vals, comments = read_labchart_text(p)
        assert len(t) == 45000 and vals.shape == (45000, 2)
        assert [c[1] for c in comments] == ["start", "15,5 inH2O", "stop", "short"]
        # csv: one header, a block column, comments in their own column
        p = tmp_path / "all.csv"
        export_selections(f, sels, [0, 1], p, fmt="csv")
        lines = p.read_text().splitlines()
        assert lines[0] == "block,time_s,Flow [V],Pmo [cmH2O],comment"
        assert lines[1].startswith("1,0,") and lines[-1].startswith("2,4.999,")
        assert sum(1 for ln in lines if ln.endswith(",start")) == 1
        # xlsx: one sheet per record
        pytest.importorskip("openpyxl")
        p = tmp_path / "all.xlsx"
        n = export_selections(f, sels, [0, 1], p, fmt="xlsx")
        assert n == 45000
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True)
        assert wb.sheetnames == ["Block 1", "Block 2"]
        ws = wb["Block 2"]
        rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
        assert rows[0][:3] == ("time_s", "Flow [V]", "Pmo [cmH2O]") and rows[0][3] == "comment"      # the first row is the titles
        assert rows[1][0] == 0.0 and abs(rows[2][0] - 0.001) < 1e-9
        assert rows[0][5:9] == ("Block", "start", "selection start (s)", "interval (s)") and rows[1][5] == 2
        # a selection longer than a worksheet is refused up front
        rec = f.records[0]
        with pytest.raises(ValueError, match="rows do not fit"):
            _refuse(f, XLSX_MAX_ROWS)
        # two selections of the same record get distinct sheet names
        p = tmp_path / "two.xlsx"
        export_selections(f, [rec.selection(0, 1), rec.selection(2, 3)], [0], p, fmt="xlsx")
        assert load_workbook(p, read_only=True).sheetnames == ["Block 1", "Block 1 (2)"]


def _refuse(f, max_rows):
    """A fake record longer than a sheet, checked without decoding anything."""
    from adichtion.export import write_xlsx
    from adichtion.file import Selection
    rec = f.records[0]
    big = Selection(0, 0, max_rows + 5)
    rec.info.ticks, saved = big.stop, rec.info.ticks     # ticks are what the row count is derived from
    try:
        write_xlsx(f, [big], [0], "unused.xlsx")
    finally:
        rec.info.ticks = saved
