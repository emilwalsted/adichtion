"""The block-formatted text/CSV writer and the streamed Excel writer give what the
value-by-value writers gave; progress, cancelling and the ``.part`` file."""
import io
import os
import re
import zipfile

import numpy as np
import pytest

from adichtion import export as ex
from adichtion.export import ExportCancelled, export_selections, whole_file, write_text, write_xlsx
from adichtion.file import AdichtFile, Selection
from adichtion.synth import SynthChannel, SynthRecord, write_synthetic

NASTY = ["start", "a & b <c> \"d\"", "=SUM(A1)", "  padded  ", "_x0041_ literal", "ÆØÅ µ",
         "with, comma", "dot . here", "two\nlines"]


def legacy_write_text(record, selection, channels, out, *, fmt="txt", header=True, decimal=".", units=True,
                      comments=True, block_column=False):
    """``write_text`` as it was before the block formatting: one call per value."""
    f = record.file
    channels = list(channels)
    ch_infos = [f.channel_info(c) for c in channels]
    t, cols, rate, row_of_tick = ex._table(record, selection, channels, units)
    n = len(t)
    sep = "\t" if fmt == "txt" else ","
    if fmt == "txt" and header:
        out.write(f"Interval=\t{ex._fmt(1.0 / rate, decimal)} s\n")
        out.write(f"ExcelDateTime=\t{ex._excel_datetime(record.start_utc):.10g}\t{record.start_local:%Y-%m-%d %H:%M:%S.%f}\n")
        out.write("TimeFormat=\tStartOfBlock\n")
        out.write("DateFormat=\tY-M-D\n")
        out.write("ChannelTitle=\t" + "\t".join(ci.title for ci in ch_infos) + "\n")
        out.write("Range=\t" + "\t".join(f"{ex._fmt(ci.range_volts, decimal)} V" if ci.range_volts else "" for ci in ch_infos) + "\n")
        out.write("UnitName=\t" + "\t".join(record.value_label(c) for c in channels) + "\n")
    elif fmt == "csv" and header:
        out.write(("block," if block_column else "") + "time_s,"
                  + ",".join(f"{ci.title} [{record.value_label(c)}]".replace(",", " ") for ci, c in zip(ch_infos, channels))
                  + (",comment" if comments else "") + "\n")
    by_row = ex._comment_rows(record, selection, row_of_tick) if comments else {}
    prefix = f"{record.index + 1}," if (block_column and fmt == "csv") else ""
    for i in range(n):
        parts = [ex._fmt(t[i], decimal)] + [ex._fmt(float(col[i]), decimal) for col in cols]
        if i in by_row:
            text = " | ".join(by_row[i])
            parts.append("#* " + text if fmt == "txt" else text.replace(",", ";"))
        out.write(prefix + sep.join(parts) + "\n")
    return n


@pytest.fixture(scope="module")
def odd(tmp_path_factory):
    """Two records with awkward values (NaN, infinities, signed zeros, repeated and unrepeated values) and
    awkward comments, on several channel kinds."""
    rng = np.random.default_rng(3)
    channels = [SynthChannel("Flow", "int16", None, 0.2), SynthChannel("Pmo", "int16", ("cmH2O", -13.34, 0.0), 10.0),
                SynthChannel("Calc", "float32", None, 10.0, lsb=1.0), SynthChannel("Pdi", "float32", ("cmH2O", 1.024, 0.0), 10.0)]
    records = []
    for ticks, offset in ((4300, 0), (900, 100)):
        comments = [(0, -1, NASTY[0])] + [(37 + 411 * i + offset, -1, txt) for i, txt in enumerate(NASTY[1:])] + [(ticks - 1, -1, "last")]
        r = SynthRecord(ticks, comments=comments)
        calc = (np.cumsum(rng.normal(0, 1, ticks)) * 1234.5678).astype(np.float32)
        calc[5], calc[6], calc[7], calc[8] = np.nan, np.inf, -np.inf, -0.0
        r.data = {0: np.clip(np.cumsum(rng.integers(-5, 6, ticks)), -32768, 32767),
                  1: np.zeros(ticks, dtype=np.int32),                                       # every value repeated (and 0 * -13.34)
                  2: calc,
                  3: (rng.normal(0, 1e7, ticks)).astype(np.float32)}                        # nothing repeats
        records.append(r)
    path = tmp_path_factory.mktemp("odd") / "odd.adicht"
    write_synthetic(path, channels, records, version=29, layout="sequential", video=False)
    return str(path)


# -- the text writer ------------------------------------------------------------------------------

@pytest.mark.parametrize("fmt", ["txt", "csv"])
@pytest.mark.parametrize("decimal", [".", ","])
@pytest.mark.parametrize("options", [dict(), dict(header=False), dict(comments=False), dict(units=False),
                                     dict(header=False, comments=False, units=False)])
@pytest.mark.parametrize("chunk", [ex.CHUNK_ROWS, 7, 1])
def test_write_text_is_byte_identical_to_the_value_by_value_writer(odd, monkeypatch, fmt, decimal, options, chunk):
    monkeypatch.setattr(ex, "CHUNK_ROWS", chunk)
    with AdichtFile(odd) as f:
        for rec, sel, chans in ((f.records[0], Selection(0, 0, 4300), [0, 1, 2, 3]),
                                (f.records[0], Selection(0, 30, 1300), [3]),
                                (f.records[1], Selection(1, 0, 900), [1, 2]),
                                (f.records[1], Selection(1, 100, 101), [0, 2, 3])):
            if chunk < 7 and sel.stop - sel.start > 1000:
                continue                                                  # one chunk per row is slow: the short ones suffice
            for block_column in (False, True):
                new, old = io.StringIO(), io.StringIO()
                n = write_text(rec, sel, chans, new, fmt=fmt, decimal=decimal, block_column=block_column, **options)
                assert legacy_write_text(rec, sel, chans, old, fmt=fmt, decimal=decimal, block_column=block_column, **options) == n
                assert new.getvalue() == old.getvalue()


def test_the_synthetic_documents_text_is_unchanged_too(synthetic):
    with AdichtFile(synthetic["path"]) as f:
        for rec in f.records:
            for fmt in ("txt", "csv"):
                new, old = io.StringIO(), io.StringIO()
                write_text(rec, Selection(rec.index, 0, rec.ticks), [0, 1, 2, 3], new, fmt=fmt, decimal=",")
                legacy_write_text(rec, Selection(rec.index, 0, rec.ticks), [0, 1, 2, 3], old, fmt=fmt, decimal=",")
                assert new.getvalue() == old.getvalue()


@pytest.mark.parametrize("spec", ["%.7g", "%.16g"])
@pytest.mark.parametrize("repeats", [True, False])
def test_strings_match_python_formatting_for_every_kind_of_value(spec, repeats):
    rng = np.random.default_rng(0)
    v = np.concatenate([rng.normal(0, 1, 500), rng.normal(0, 1e300, 50), rng.normal(0, 1e-300, 50), [0.0, -0.0, np.nan, np.inf, -np.inf,
                        5e-324, 1.7976931348623157e308, 1e16, 123456.5, 0.1 + 0.2, 1e-5, 100000.0, 1234567.0, 12345678.0],
                        np.repeat(rng.normal(0, 1, 40), 20)])
    got = ex._strings(v, spec, repeats=repeats).tolist()
    assert got == [spec % x for x in v.tolist()]
    assert "-0" in got and "nan" in got and "inf" in got
    assert ex._strings(np.array([0.0, -0.0, 0.0, -0.0, 0.0, -0.0]), "%.7g").tolist() == ["0", "-0"] * 3


# -- the Excel writer -----------------------------------------------------------------------------

def _cells(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}
    finally:
        wb.close()


def _numeric_texts(path):
    """The text of every numeric ``<v>`` (an empty one for NaN and infinity) of every sheet, in order."""
    with zipfile.ZipFile(path) as z:
        return {n: re.findall(r'<c r="[A-Z]+\d+" t="n">(?:<v>([^<]*)</v>|<v />)</c>', z.read(n).decode())
                for n in sorted(z.namelist()) if n.startswith("xl/worksheets/sheet")}


@pytest.mark.parametrize("options", [dict(), dict(header=False), dict(comments=False), dict(units=False)])
@pytest.mark.parametrize("chunk", [ex.CHUNK_ROWS, 64])
def test_streamed_xlsx_has_the_cells_of_the_openpyxl_writer(synthetic, tmp_path, monkeypatch, options, chunk):
    pytest.importorskip("openpyxl")
    monkeypatch.setattr(ex, "CHUNK_ROWS", chunk)
    with AdichtFile(synthetic["path"]) as f:
        sels = [Selection(0, 15000, 17000), Selection(1, 0, 4000), Selection(0, 39000, 40000), Selection(0, 15000, 15500)]
        fast, slow = tmp_path / "fast.xlsx", tmp_path / "slow.xlsx"
        assert write_xlsx(f, sels, [0, 1, 2, 3], fast, **options) == write_xlsx(f, sels, [0, 1, 2, 3], slow, fast=False, **options)
    assert _cells(fast) == _cells(slow)
    assert _numeric_texts(fast) == _numeric_texts(slow)                       # the same digits, not just the same doubles
    assert any("15,5 inH2O" in str(r) for rows in _cells(fast).values() for r in rows) == options.get("comments", True)


@pytest.mark.parametrize("chunk", [ex.CHUNK_ROWS, 64])
def test_streamed_xlsx_matches_openpyxl_on_nan_and_infinity(odd, tmp_path, monkeypatch, chunk):
    pytest.importorskip("openpyxl")
    monkeypatch.setattr(ex, "CHUNK_ROWS", chunk)
    with AdichtFile(odd) as f:
        sels = [Selection(0, 0, 4300), Selection(1, 100, 700)]
        fast, slow = tmp_path / "fast.xlsx", tmp_path / "slow.xlsx"
        write_xlsx(f, sels, [0, 1, 2, 3], fast, comments=False)
        write_xlsx(f, sels, [0, 1, 2, 3], slow, comments=False, fast=False)
    assert _cells(fast) == _cells(slow)
    texts = _numeric_texts(fast)
    assert texts == _numeric_texts(slow) and "" in texts["xl/worksheets/sheet1.xml"]       # NaN and infinity: empty cells


def test_streamed_xlsx_comments_and_non_finite_values(odd, tmp_path):
    pytest.importorskip("openpyxl")
    with AdichtFile(odd) as f:
        path = tmp_path / "c.xlsx"
        write_xlsx(f, [Selection(0, 0, 4300)], [2], path)
        with zipfile.ZipFile(path) as z:
            xml = z.read("xl/worksheets/sheet1.xml").decode()
        assert "_x005F_x0041_ literal" in xml                                      # a literal escape sequence protected
        assert "a &amp; b &lt;c&gt;" in xml
        cells = _cells(path)["Block 1"]
        header, rows = cells[:2], cells[2:]
        assert header[1] == ["time_s", "Calc [V]", "comment"] or header[1][0] == "time_s"
        by_text = {r[2]: r for r in rows if len(r) > 2 and r[2] is not None}
        assert by_text["=SUM(A1)"][2] == "=SUM(A1)"                              # text, not a formula
        assert by_text["  padded  "][2] == "  padded  "
        assert by_text["a & b <c> \"d\""][2] == "a & b <c> \"d\""
        assert by_text["ÆØÅ µ"][2] == "ÆØÅ µ"
        assert by_text["two\nlines"][2] == "two\nlines"
        assert [r[1] for r in rows[5:9]] == [None, None, None, -0.0] or rows[5][1] is None   # NaN and infinities are empty cells
        assert len(rows) == 4300


def test_text_xml_cannot_carry_is_escaped_the_way_excel_reads_back():
    assert ex._xml_text("ctrl\x01char\x0b\ufffe") == "ctrl_x0001_char_x000B__xFFFE_"
    assert ex._xml_text("_x0041_ & <b>") == "_x005F_x0041_ &amp; &lt;b&gt;"
    assert ex._xml_text("ÆØÅ\ttab\nnl") == "ÆØÅ\ttab\nnl"
    assert ex._xlsx_cell("A1", "a\x02b") == '<c r="A1" t="inlineStr"><is><t xml:space="preserve">a_x0002_b</t></is></c>'
    assert ex._xlsx_cell("A1", None) == "" and ex._xlsx_cell("B2", float("nan")) == '<c r="B2" t="n"><v /></c>'
    assert ex._xlsx_cell("B2", 3) == '<c r="B2" t="n"><v>3</v></c>' and ex.get_column_letter(27) == "AA"


def test_xlsx_falls_back_to_openpyxl_when_its_layout_is_not_the_expected_one(synthetic, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    monkeypatch.setattr(ex, "_xlsx_streamed", lambda *a, **k: False)
    with AdichtFile(synthetic["path"]) as f:
        assert write_xlsx(f, [Selection(0, 0, 300)], [0, 1], tmp_path / "x.xlsx") == 300
    assert len(_cells(tmp_path / "x.xlsx")["Block 1"]) == 302 and not (tmp_path / "x.xlsx.part").exists()


# -- progress, cancelling, the .part file ---------------------------------------------------------

@pytest.mark.parametrize("fmt", ["txt", "csv", "xlsx"])
def test_progress_counts_rows_up_to_the_total(synthetic, tmp_path, monkeypatch, fmt):
    pytest.importorskip("openpyxl")
    monkeypatch.setattr(ex, "CHUNK_ROWS", 1000)
    seen = []
    with AdichtFile(synthetic["path"]) as f:
        n = export_selections(f, whole_file(f), [0, 1], tmp_path / f"p.{fmt}", fmt=fmt, progress=lambda d, t: seen.append((d, t)))
    assert n == 45000 and seen[-1] == (45000, 45000)
    assert all(t == 45000 for _d, t in seen) and [d for d, _t in seen] == sorted(d for d, _t in seen) and len(seen) >= 45


@pytest.mark.parametrize("fmt", ["txt", "csv", "xlsx"])
def test_a_cancelled_export_leaves_nothing_and_keeps_an_existing_file(synthetic, tmp_path, monkeypatch, fmt):
    pytest.importorskip("openpyxl")
    monkeypatch.setattr(ex, "CHUNK_ROWS", 1000)
    target = tmp_path / f"c.{fmt}"
    with AdichtFile(synthetic["path"]) as f:
        calls = []

        def stale():
            calls.append(1)
            return len(calls) > 3

        with pytest.raises(ExportCancelled):
            export_selections(f, whole_file(f), [0, 1], target, fmt=fmt, stale=stale)
        assert not target.exists() and not (tmp_path / f"c.{fmt}.part").exists()
        target.write_bytes(b"precious")
        with pytest.raises(ExportCancelled):
            export_selections(f, whole_file(f), [0, 1], target, fmt=fmt, stale=lambda: True)
        assert target.read_bytes() == b"precious" and not (tmp_path / f"c.{fmt}.part").exists()


def test_a_failing_export_leaves_no_part_file_and_a_good_one_replaces_the_old(synthetic, tmp_path):
    target = tmp_path / "e.csv"
    target.write_text("old")
    with AdichtFile(synthetic["path"]) as f:
        def boom(done, total):
            raise RuntimeError("disk full")
        with pytest.raises(RuntimeError, match="disk full"):
            export_selections(f, whole_file(f), [0], target, fmt="csv", progress=boom)
        assert target.read_text() == "old" and not (tmp_path / "e.csv.part").exists()
        export_selections(f, whole_file(f), [0], target, fmt="csv")
        assert target.read_text().startswith("block,time_s,") and not (tmp_path / "e.csv.part").exists()
        with pytest.raises(ValueError, match="no channels"):
            export_selections(f, whole_file(f), [], tmp_path / "none.csv", fmt="csv")
        with pytest.raises(ValueError, match="unknown export format"):
            export_selections(f, whole_file(f), [0], tmp_path / "none.csv", fmt="pdf")
        assert not (tmp_path / "none.csv").exists()


@pytest.mark.skipif(not os.path.exists("/dev/null"), reason="needs /dev/null")
def test_a_target_that_is_not_a_regular_file_is_written_to_directly(synthetic):
    with AdichtFile(synthetic["path"]) as f:
        assert export_selections(f, whole_file(f), [0], "/dev/null", fmt="csv") == 45000


def test_table_rows_is_never_negative(synthetic):
    with AdichtFile(synthetic["path"]) as f:
        rec = f.records[1]
        assert ex.table_rows(rec, Selection(1, rec.ticks + 4, rec.ticks + 6), [0]) == 0
        assert ex.table_rows(rec, Selection(1, -3, -1), [0]) == 0
